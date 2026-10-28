from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "deliverables"
QA_DIR = ROOT / "qa"
OUT_DIR.mkdir(exist_ok=True)
QA_DIR.mkdir(exist_ok=True)
DIAGRAM = QA_DIR / "backend-architecture.png"
OUTPUT = OUT_DIR / "HAN2601_Planning_Report.docx"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EFF6FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9D9D9"
TEXT = "1F2937"
MUTED = "5B6573"
WHITE = "FFFFFF"


def font(size, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "",
        "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    if bold:
        candidates = [
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
            "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ]
    for p in candidates:
        if p and Path(p).exists():
            return ImageFont.truetype(p, size=size)
    return ImageFont.load_default()


def rounded_box(draw, xy, fill, outline, title, subtitle=None, title_fill="#17365D", title_size=27, subtitle_size=18):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
    tf = font(title_size, True)
    sf = font(subtitle_size, False)
    if subtitle:
        ty = y1 + 17
        draw.text(((x1+x2)/2, ty), title, font=tf, fill=title_fill, anchor="ma")
        draw.text(((x1+x2)/2, ty+39), subtitle, font=sf, fill="#4B5563", anchor="ma")
    else:
        draw.text(((x1+x2)/2, (y1+y2)/2), title, font=tf, fill=title_fill, anchor="mm")


def arrow(draw, start, end, color="#2F75B5", width=5):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1]-start[1], end[0]-start[0])
    length = 18
    pts = [
        end,
        (end[0] - length*math.cos(angle-0.55), end[1] - length*math.sin(angle-0.55)),
        (end[0] - length*math.cos(angle+0.55), end[1] - length*math.sin(angle+0.55)),
    ]
    draw.polygon(pts, fill=color)


def make_diagram():
    img = Image.new("RGB", (1600, 820), "white")
    d = ImageDraw.Draw(img)
    rounded_box(d, (500, 25, 1100, 118), "#EFF6FB", "#2F75B5", "Web interface and experiment client", "REST commands and WebSocket event stream")
    rounded_box(d, (540, 165, 1060, 260), "#D9EAF7", "#2F75B5", "API and session service", "configuration, lifecycle and timeouts")
    arrow(d, (800, 118), (800, 165))

    rounded_box(d, (55, 322, 605, 470), "#F2F4F7", "#7B8794", "Agent gateway", "one Observe and Act contract")
    d.rounded_rectangle((650, 305, 1545, 675), radius=16, fill="#EFF6FB", outline="#17365D", width=3)
    d.text((1098, 330), "Deterministic auction engine", font=font(29, True), fill="#17365D", anchor="ma")
    d.text((1098, 370), "authoritative state machine and rule enforcement", font=font(19), fill="#4B5563", anchor="ma")
    rounded_box(d, (685, 430, 950, 620), "#FFFFFF", "#7B8794", "Mechanism plugins", "required auction types", title_size=23, subtitle_size=17)
    rounded_box(d, (970, 430, 1240, 620), "#FFFFFF", "#7B8794", "Settlement", "winner, payment and utility", title_size=23, subtitle_size=17)
    rounded_box(d, (1260, 430, 1510, 620), "#FFFFFF", "#7B8794", "Action validator", "schema, budget and legal move", title_size=22, subtitle_size=16)
    rounded_box(d, (80, 515, 325, 650), "#FFFFFF", "#7B8794", "Programmed agent", "heuristic or custom code", title_size=22, subtitle_size=16)
    rounded_box(d, (350, 515, 580, 650), "#FFFFFF", "#7B8794", "LLM agent", "prompt, schema and retries", title_size=23, subtitle_size=16)
    rounded_box(d, (710, 715, 1485, 805), "#F2F4F7", "#7B8794", "Event log and experiment store", "replay, metrics, prompts, costs and results", title_size=25, subtitle_size=17)

    arrow(d, (800, 260), (800, 305))
    arrow(d, (650, 400), (605, 400))
    arrow(d, (605, 435), (650, 435))
    arrow(d, (202, 470), (202, 515), color="#7B8794")
    arrow(d, (465, 470), (465, 515), color="#7B8794")
    arrow(d, (1098, 675), (1098, 715))
    img.save(DIAGRAM, dpi=(220, 220))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_run_font(run, name="Arial", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, before=0, after=4.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.06
    p.add_run(text)
    return p


def add_bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.03
    r = p.add_run(lead)
    r.bold = True
    p.add_run(text)
    return p


def build_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(4.5)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(3)
    title_ppr = title._element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)

    for name, size, before, after in (("Heading 1", 12.2, 6, 2), ("Heading 2", 10.8, 4, 1.5)):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Arial"
        styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        styles[name].font.size = Pt(9.8)

    props = doc.core_properties
    props.title = "Planning Report for Generic Auction Simulator for Programmed and LLM Based Agents"
    props.subject = "Final Year Project HAN2601"
    props.author = "LIU Xianlong and MA Ho Tsun Matthew"
    props.keywords = "auction simulator, LLM agents, backend architecture, final year project"

    p = doc.add_paragraph(style="Title")
    p.add_run("Planning Report for a Generic Auction Simulator")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(5)
    r = sub.add_run("Programmed and LLM Based Agents")
    set_run_font(r, size=12.4, bold=True, color=NAVY)

    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    widths = [0.72, 2.25, 0.85, 3.18]
    for row in meta.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            set_cell_margins(row.cells[i], 55, 75, 55, 75)
            row.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    entries = [
        ("Code", "HAN2601", "Supervisor", "Prof. ZHANG Hanrui"),
        ("Students", "LIU Xianlong and MA Ho Tsun Matthew", "Submission", "30 September 2026"),
    ]
    for ridx, vals in enumerate(entries):
        for cidx, val in enumerate(vals):
            cell = meta.rows[ridx].cells[cidx]
            cell.text = val
            if cidx % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for rr in cell.paragraphs[0].runs:
                    set_run_font(rr, size=8.8, bold=True, color=NAVY)
            else:
                for rr in cell.paragraphs[0].runs:
                    set_run_font(rr, size=8.8, color=TEXT)
    set_table_borders(meta)

    add_heading(doc, "Project Goal")
    add_body(doc, "The project will deliver a web based auction simulator in which programmed bidders and large language model agents participate through the same interface. The simulator will run sealed bid mechanisms, including first price, second price and all pay auctions, and dynamic mechanisms, including English, Dutch and Japanese auctions. It will expose each bidder only to permitted observations, apply valuation and budget constraints, determine allocations and payments, and record a complete event history for replay and analysis. Multi item combinatorial auctions are a stretch goal after the core mechanisms are validated.")

    add_heading(doc, "Significance of the Project")
    add_body(doc, "Auction strategy is difficult to study when every mechanism, agent and experiment uses different code. This project will provide a common environment for controlled comparisons. Students and researchers will be able to test hand coded strategies against LLM based strategies under identical valuations, budgets and auction rules. The visual replay will also help lecturers and learners inspect how bids, beliefs and outcomes change over time. The project matters because dynamic competitive tasks test planning, adaptation, rule following and resource management more directly than a static question answering benchmark.")

    add_heading(doc, "Problem Statement")
    add_body(doc, "A useful simulator must support different auction rules without rewriting the whole system, while keeping comparisons fair and reproducible. LLM agents add further uncertainty: they may return malformed actions, exceed response deadlines, violate budgets or vary between runs. If the auctioneer also relies on free form LLM output, an invalid decision can change the result. The backend therefore needs a deterministic source of truth for rules and state, a uniform agent contract, strict action validation, fault isolation, seeded experiments, and an auditable record of every observation, action and outcome.")

    add_heading(doc, "Proposed Solution and Deliverables")
    add_bullet(doc, "Core simulation backend. ", "A typed domain model and deterministic state machine for auction sessions, rounds, lots, bidders, valuations, bids, allocations and payments.")
    add_bullet(doc, "Extensible mechanism library. ", "A plug in contract for auction rules, followed by the three sealed bid and three dynamic mechanisms in the project brief.")
    add_bullet(doc, "Common agent interface. ", "Adapters for custom Python strategies and LLM providers using the same Observe and Act messages, with structured outputs, timeouts, retries and budget checks.")
    add_bullet(doc, "Experiment and data services. ", "Batch runs with fixed seeds, an append only event log, prompt and cost metadata, summary metrics, export and deterministic replay.")
    add_bullet(doc, "Web demonstration and evaluation. ", "A human friendly live view and replay, automated rule tests, agent failure tests, benchmark scenarios, documentation and the final report.")

    scope = doc.add_paragraph()
    scope.paragraph_format.space_before = Pt(2)
    scope.paragraph_format.space_after = Pt(0)
    r = scope.add_run("Scope control  ")
    set_run_font(r, size=9.6, bold=True, color=NAVY)
    r = scope.add_run("The minimum viable product ends with reliable single item mechanisms, two agent types and reproducible replay. Combinatorial auctions, distributed workers and additional LLM providers will be attempted only after that baseline passes its tests.")
    set_run_font(r, size=9.6, color=TEXT)

    doc.add_page_break()

    add_heading(doc, "Backend Architecture", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DIAGRAM), width=Inches(7.0))
    cap = doc.add_paragraph("Figure 1  Proposed backend architecture and principal event flow")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(3)
    for r in cap.runs:
        set_run_font(r, size=8.2, color=MUTED)

    add_body(doc, "The auction engine, rather than an LLM, is the authority for legal state transitions, winners and payments. The session service creates runs and publishes events; the agent gateway gives every bidder the same observation and action schema; and the experiment store preserves the evidence needed to reproduce a result. An optional LLM auctioneer can later explain events in natural language, but it must not override the deterministic engine.", after=2.5)

    add_heading(doc, "Key Backend Decisions", 2)
    decisions = doc.add_table(rows=2, cols=2)
    decisions.alignment = WD_TABLE_ALIGNMENT.CENTER
    decisions.autofit = False
    decision_text = [
        ("Modular monolith first", "One deployable backend with clear modules reduces integration risk; workers can be separated later if batch experiments require it."),
        ("Event sourced sessions", "Every accepted action and state change becomes an event, enabling replay, audit and live WebSocket updates."),
        ("Deterministic boundaries", "Seeded valuations and coded auction rules make comparisons repeatable even when an LLM response is stochastic."),
        ("Provider neutral agents", "Adapters isolate prompts and vendor APIs from the auction domain, so programmed and LLM bidders share one contract."),
    ]
    for idx, (lead, desc) in enumerate(decision_text):
        cell = decisions.rows[idx // 2].cells[idx % 2]
        cell.text = ""
        set_cell_margins(cell, 75, 90, 75, 90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx in (0, 3):
            set_cell_shading(cell, PALE_BLUE)
        pp = cell.paragraphs[0]
        pp.paragraph_format.space_after = Pt(0)
        rr = pp.add_run(lead + "  ")
        set_run_font(rr, size=8.6, bold=True, color=NAVY)
        rr = pp.add_run(desc)
        set_run_font(rr, size=8.6, color=TEXT)
    set_table_borders(decisions)

    add_heading(doc, "Proposed Timeline", 1)
    timeline = doc.add_table(rows=1, cols=3)
    timeline.alignment = WD_TABLE_ALIGNMENT.CENTER
    timeline.autofit = False
    timeline.columns[0].width = Inches(1.35)
    timeline.columns[1].width = Inches(2.0)
    timeline.columns[2].width = Inches(3.65)
    hdr = timeline.rows[0]
    set_repeat_table_header(hdr)
    for i, txt in enumerate(("Period", "Work package", "Expected output")):
        c = hdr.cells[i]
        c.text = txt
        set_cell_shading(c, NAVY)
        set_cell_margins(c, 70, 90, 70, 90)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for rr in c.paragraphs[0].runs:
            set_run_font(rr, size=8.4, bold=True, color=WHITE)
    rows = [
        ("23 Sep–18 Oct 2026", "Requirements and design", "Planning report, use cases, message schemas, mechanism contract and test plan"),
        ("19 Oct–13 Dec 2026", "Auction engine", "Sealed bid and dynamic mechanisms with unit and invariant tests"),
        ("14 Dec 2026–31 Jan 2027", "Agent integration", "Programmed adapter, LLM adapter, validation, timeout and cost controls"),
        ("1–28 Feb 2027", "Experiments and storage", "Seeded batch runner, event store, replay, metrics and export"),
        ("1–31 Mar 2027", "Web demonstration", "Live event stream, configuration interface and visual replay"),
        ("1 Apr 2027–final deadline", "Evaluation and submission", "Benchmarks, reliability tests, analysis, documentation, demo and final report"),
    ]
    for ridx, vals in enumerate(rows):
        row = timeline.add_row()
        for cidx, val in enumerate(vals):
            c = row.cells[cidx]
            c.text = val
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c, 55, 75, 55, 75)
            if ridx % 2 == 1:
                set_cell_shading(c, LIGHT_GRAY)
            for rr in c.paragraphs[0].runs:
                set_run_font(rr, size=7.9, color=TEXT)
    set_table_borders(timeline)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Success criteria  ")
    set_run_font(r, size=8.6, bold=True, color=NAVY)
    r = p.add_run("Every required mechanism passes rule and payment tests; invalid agent actions cannot corrupt a session; a seeded programmed-agent experiment replays to the same result; and an LLM-agent run remains traceable with latency, token and cost records. The dates after March will be aligned with the confirmed departmental submission schedule.")
    set_run_font(r, size=8.6, color=TEXT)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Reference  ")
    set_run_font(r, size=8.0, bold=True, color=NAVY)
    r = p.add_run("Chen, J. et al. (2024), Put Your Money Where Your Mouth Is: Evaluating Strategic Planning and Execution of LLM Agents in an Auction Arena, arXiv:2310.05746; project page: https://auction-arena.github.io/.")
    set_run_font(r, size=8.0, color=MUTED)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fr = fp.add_run("HAN2601  |  Planning Report  |  30 September 2026")
    set_run_font(fr, size=8, color=MUTED)

    doc.save(OUTPUT)


if __name__ == "__main__":
    make_diagram()
    build_doc()
    print(OUTPUT)
